"""Author-facing revision response to the supplied simulated pre-submission review."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'editorial/review_20261027';S=ROOT/'supplementary/review_20261027'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def main():
 stats=read(S/'REPEATED_SEED_STATISTICS.json');d=stats['completed_transport_distributions']
 items=[
 ('本轮修改范围','本文件回应用户提供的《AST_PreSubmission_Reviewer_Report_CN.docx》，属于投稿前内部修订记录，并非向期刊提交的正式回复。英文主稿保持研究论文写法；未完成的审稿请求在本文件中逐项说明。作者署名、单位、通讯作者及基金保留原有信息。'),
 ('M1 最新研究与贡献定位：部分落实','新增并核对 Chai、Sun、Chan 的三篇 2026 年 AST 论文及 DOI，当前共 19 条参考文献。引言以段落说明任务与研究目标，按用户要求不新增综述式文献比较表。标题及贡献聚焦“惯性信息如何通过绳力约束影响支撑编队”，不再暗示从未知环境自主发现姿态序列。三篇新增文献的元数据已保存；尚未取得足以完成逐项算法比较的全文依据，也未新增其算法的同任务实现。因此本轮不宣称超过这些方法，全文核读与强规划基线仍需补充。'),
 ('M2 十维辨识与预测：已补充对应数值证据','对四种基本配置和两种叉形课程补齐真实与估计质量、三轴质心、原点惯量、质心惯量、Frobenius 相对误差及主惯量误差。质量误差为 0.0024%—0.0188%，质心惯量误差为 5.46%—7.87%。图 3 同时展示质量/惯量误差、质心分量、观测条件数和运输阶段预测差。运输阶段完整 0.4 s 动量窗口排除于准备阶段拟合之外；这属于同次执行中的留出阶段，并未表述成另一条新激励飞行。正则化从 10⁻¹⁰ 到 10⁻⁴ 的六个记录重拟合，质心误差范数变化不超过 0.038 mm，质心惯量误差变化不超过 0.092 个百分点。残差未进行协方差白化，文中交代该选择并分别提供力/力矩单位下的残差。复核时统一采用原始扫描包围盒中点坐标，保证扫描、吊点和惯性参数同框架。'),
 ('M3 公平消融：五组已完成，既有算法基线尚未加入','新增 3/4 机 × 矩形/不规则 × 五种惯性信息，共 20 个初始尝试。五组为真实质量匹配的均匀模型、仅辨识质量、辨识质量+质心、完整惯性辨识、真实惯性理想参考。参考、候选预算、增益及配重保持相同。四机矩形的姿态峰值依次约为 29.66°、29.44°、0.731°、0.712°、0.678°，因此明确将该任务的主要收益归于质心修正，未将全部改善归因于转动惯量。表格分别标注规划拒绝、采集超时和完成执行超限。三机不规则质量+质心组首次在起运前因图像年龄超时中止，另行保留同参数采集重试，原尝试未删除或替换。真实惯性组被称为信息参考，未冒充另一种规划算法。'),
 ('M4 自动决策范围：已统一','给定输入为障碍地图、平滑载荷位姿参考族、标定吊点；连续决策为各节点绳力，经顺应性映射得到支撑端点，有限搜索选择时长和下沉量。标题、摘要、图 2、规划节及结论均采用同一范围。已实现的是起运前整段力/几何规划及运动中的闭环反馈，不称为未知树林在线姿态序列搜索。'),
 ('M5 重复与不确定性：部分落实',f"新增四机矩形固定任务的 20 个独立传感种子初始尝试：{stats['passed']} 次完成并通过、{stats['acquisition_aborts']} 次触发图像年龄门限，全部保留。正文报告全尝试率 {stats['passed']}/20，Wilson 95% 区间 {stats['wilson95_all_attempts'][0]*100:.1f}%—{stats['wilson95_all_attempts'][1]*100:.1f}%，并将完成任务的 RMS/峰值分布明确标成条件统计。完成运输的位置 RMS 中位数/p95 为 {d['position_rms_cm']['median']:.2f}/{d['position_rms_cm']['p95']:.2f} cm，姿态 RMS 为 {d['attitude_rms_deg']['median']:.3f}/{d['attitude_rms_deg']['p95']:.3f}°。另对两个封存图像位姿记录进行 160 次观测层重拟合，覆盖附加位姿噪声、张力噪声/零偏、固定姿态框架偏差和随机丢帧；扰动范围作为声明的合成敏感性参数，未冒充传感器实测规格。尚未覆盖每种构型各 20—30 次、像素级遮挡、时间戳失配、缆索和电机参数失配的全矩阵实验。当前 20 次也反映了真实墙钟感知链的中止率，不作为普遍鲁棒性结论。"),
 ('M6 三级可行性与裕度：表述已修正，参数敏感性未闭合','正文分别列出锥优化节点约束、非线性候选复核、闭环物理执行门限。澄清线性化参考间距 0.72 m、候选接纳间距 0.52 m、实际机中心最小距离 0.640 m 与扣除包络后的障碍净空 0.081 m 的区别。说明名义张力、动态推力和单电机上限，并区分准备全段极值与运输统计。20 ms 几何检查、完整绳段、末端宽度及半采样运动膨胀均记录；这些是所测轨迹的采样几何裕度，不宣称任意扰动下连续时间安全定理。分配权重、d_lin 和迭代轮数的系统敏感性尚未加入，未将此项标成全部完成。'),
 ('M7 独立复核与开放复现：已扩展数值层级','独立 NumPy/RK4 电机命令回放从一个基本区间扩展到大幅滚转、张力变化最大和飞机包络净空最小的三个课程区间。每段持续 2 s，内部无状态重置；均比较 1、0.5、0.25 ms MuJoCo 步长。在 0.5 ms 下，新增三个区间载荷位置差不超过 0.365 mm、姿态差不超过 0.040°。图 9 将这些结果与处理时间经验分布共同展示。公开代码/数据目标仓库为 https://github.com/shepherd-p/ForestSling；最终发布状态以 PUBLICATION_CHECK.json 为准。MuJoCo 闭环、Unreal 记录渲染与独立力学回放的用途均明确；没有增加实物飞行，也没有声称独立实物或真实感知链验证。'),
 ('m1—m3 仿真、延迟与非凸支撑','摘要明确所有定量结果来自仿真；准备阶段九帧微分产生的 0.4 s 固定滞后与执行时传感年龄分开解释。非凸扫描凸包仅作为惯性矩外包络，碰撞仍使用非凸表面；新增只保留评估器实体材料内支撑点的事后诊断，两叉形记录的质心惯量误差为 5.85% 和 6.72%，原凸包为 5.86% 和 6.77%。诊断使用的评估器几何不进入在线辨识。'),
 ('m4—m8 参数、符号和排版','补充求解器容差和迭代上限，原始候选保留实际迭代数；目标权重、λ、ε 与质量区间均在方法中列出。补充 25 ms 电机常数、0.2 N s/m 阻力系数、ICP 阈值、绳端长度与受力位置。统一重力标量/向量、标题连字符、载荷重定向与编队重构的定义。图 3 明确误差含义；图 8 明确机中心间距。未用没有依据的协方差或传感器规格解释参数。'),
 ('用户要求 1—6：图表与数据可用性','图 1 采用单一空间路线示意与三处原始 Unreal 通过视图，移除测得滚转子图。图 2 采用扫描—辨识—重构的机械信息流，不再把多张照片拼成框图。图 3 改为完整惯性误差图谱、质心分量、可观测性及留出预测；图 9 改为独立复核误差图谱、视觉时长经验分布与控制时间预算。所有九图统一 Arial 字体与配色，数学标注采用 STIX；正文表格采用 Arial。PDF 与 SVG 为可编辑矢量源，600 dpi 框架预览保留。数据可用性改为明确的公开 GitHub 地址，不再要求读者仅凭一个看不到的 S1 名称寻找数据。'),
 ('本轮后仍需作者决策的投稿证据','尚未闭合的主要请求集中为：直接相关现有算法的同任务基线；跨更多构型及真实工程范围的感知/模型扰动重复；规划权重和间距线性化的参数敏感性。它们已经在本回复中标明，未伪装成已完成结果。当前交付是一轮有新增实验和公开复现材料的研究稿修订，不代表期刊接收承诺。')]
 text='# AST 投稿前审阅意见逐项回复\n\n修订日期：2026-10-27\n\n'+'\n\n'.join('## '+h+'\n\n'+t for h,t in items)
 (OUT/'Revision_Response_CN.md').write_text(text+'\n',encoding='utf-8')
 doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(10.5 if name=='Normal' else 13)
  st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei' if name!='Normal' else 'SimSun')
  st.font.color.rgb=RGBColor.from_string('203A49' if name!='Normal' else '000000')
  st.paragraph_format.space_after=Pt(7);st.paragraph_format.widow_control=True
 doc.add_heading('AST 投稿前审阅意见逐项回复',0);doc.add_paragraph('修订日期：2026-10-27')
 for h,t in items:doc.add_heading(h,1);p=doc.add_paragraph(t);p.paragraph_format.line_spacing=1.2
 footer=sec.footer.paragraphs[0];footer.alignment=1;field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.core_properties.title='AST 投稿前审阅意见逐项回复';doc.core_properties.author='Fangwei Peng and coauthors'
 doc.save(OUT/'Revision_Response_CN.docx')
 (OUT/'REVIEW_STATUS.json').write_text(json.dumps(dict(revision_date='2026-10-27',completed=['M2 numerical evidence','M4 scope','M7 expanded numerical replay','figure redesign','font unification'],partial=['M1 latest-method baseline','M3 existing-method baseline','M5 broad perturbation coverage','M6 parameter sensitivity'],all_comments_closed=False),ensure_ascii=False,indent=2),encoding='utf-8')
 print('Written revision response')
if __name__=='__main__':main()

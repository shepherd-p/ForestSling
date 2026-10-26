"""Create the Chinese manuscript with editable paragraphs, tables and equations."""
from pathlib import Path
import json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1];M=ROOT/'manuscript'
E=json.loads((ROOT/'supplementary/evidence.json').read_text(encoding='utf-8'))
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
sec.footer_distance=Cm(.9)

def font_style(style,size,east='SimSun',latin='Times New Roman'):
    style.font.name=latin;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts();fonts.set(qn('w:eastAsia'),east)

font_style(doc.styles['Normal'],10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.2
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.widow_control=True
for name,size in [('Title',16),('Heading 1',12),('Heading 2',11)]:
    font_style(doc.styles[name],size,'Microsoft YaHei')
    doc.styles[name].paragraph_format.space_before=Pt(10 if name!='Title' else 0)
    doc.styles[name].paragraph_format.space_after=Pt(6)
    doc.styles[name].paragraph_format.keep_with_next=True
font_style(doc.styles['Caption'],9)
doc.styles['Caption'].paragraph_format.line_spacing=1.1
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].font.bold=False
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)

def paragraph(text='',style=None):
    p=doc.add_paragraph(text,style);p.paragraph_format.widow_control=True
    return p

def author_block():
    p=paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for i,(name,unit) in enumerate([('Fangwei Peng','a'),('Juntong Qi','b,*'),('Yan Peng','b'),('Mingming Wang','c'),('Chong Wu','d')]):
        if i:p.add_run(', ')
        p.add_run(name);r=p.add_run(unit);r.font.superscript=True
    addresses=[
      'a School of Mechatronic Engineering and Automation, Shanghai University, Shanghai 200444, China',
      'b School of Future Technology, Shanghai University, Shanghai 200444, China',
      'c School of Electrical and Information Engineering, Tianjin University, Tianjin 300072, China',
      'd EFY Intelligent Control (Tianjin) Technology Group Co., Ltd., Tianjin 300450, China',
      '* Corresponding author: Juntong Qi, qijt@shu.edu.cn']
    for address in addresses:
        p=paragraph(address);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.0
        for r in p.runs:r.font.size=Pt(8.5)

def equation(text):
    p=paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(7)
    node=OxmlElement('m:oMath');run=OxmlElement('m:r')
    prop=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
    for key in ('ascii','hAnsi','eastAsia'):fonts.set(qn('w:'+key),'Cambria Math')
    prop.append(fonts);size=OxmlElement('w:sz');size.set(qn('w:val'),'21');prop.append(size)
    run.append(prop);t=OxmlElement('m:t');t.text=text;run.append(t);node.append(run);p._p.append(node)

def table(caption,headers,rows,widths):
    p=paragraph(caption,'Caption');p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    for cell,width,text in zip(t.rows[0].cells,widths,headers):cell.width=Cm(width);cell.text=text
    for row in rows:
        cells=t.add_row().cells
        for cell,width,text in zip(cells,widths,row):cell.width=Cm(width);cell.text=str(text)
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','bottom','left','right'):
                item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
            props.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','bottom','left','right'):
                item=OxmlElement('w:'+edge);item.set(qn('w:w'),'65');item.set(qn('w:type'),'dxa');margins.append(item)
            props.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE5EB');props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
                p.paragraph_format.keep_with_next=i<len(t.rows)-1
                for r in p.runs:r.font.name='Arial';r.font.size=Pt(8.5 if len(headers)>4 else 9);r.bold=i==0
        pr=row._tr.get_or_add_trPr();dont=OxmlElement('w:cantSplit');pr.append(dont)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    paragraph().paragraph_format.space_after=Pt(0)

def add_table(key):
    rows=sorted([r for r in E['matrix'] if r['method']=='estimated'],key=lambda r:(r['count'],r['shape']=='irregular'))
    shapes={'rectangular':'矩形','irregular':'不规则'}
    if key=='parameters':
        table('表 1 动力学及感知参数',['参数','数值'],[
          ['无人机质量','1.0 kg'],['机体惯量对角元素','[0.009, 0.009, 0.016] kg m²'],['缆索静长','2.0 m'],
          ['刚度／轴向阻尼','180 N/m／5 N s/m'],['单电机推力上限','7.5 N'],['控制／图像请求','50 Hz／10 Hz'],
          ['电机时间常数／线性阻力','25 ms／0.2 N s/m'],['ICP RMSE／内点比例门限','18 mm／0.65'],['对应距离门限','80 mm'],
          ['每视角 RGB-D 分辨率','256×192'],['深度噪声标准差','1 mm'],['张力噪声','均匀 ±0.01 N'],
          ['无人机位置噪声','均匀 ±0.5 mm'],['辨识准备运动','16 s'],['物理积分步长','0.5 ms']],[6.5,10.1])
    elif key=='matrix':
        table('表 2 四种辨识配置的运输结果',['机数','外形','运输\n(s)','质心误差\n(mm)','位置峰值\n(cm)','姿态峰值\n(°)','张力\n(N)'],[
            [r['count'],shapes[r['shape']],f"{r['planned_motion_duration_sec']:.0f}",f"{r['com_error_m']*1000:.2f}",f"{r['max_position_error_m']*100:.2f}",f"{r['max_attitude_error_deg']:.2f}",f"{r['min_tension_n']:.2f}–{r['max_tension_n']:.2f}"] for r in rows],[1.,1.7,1.5,2.8,2.8,2.8,4.])
    elif key=='baseline':
        rr=[]
        for r in rows:
            u=next(p for p in E['matrix'] if p['count']==r['count'] and p['shape']==r['shape'] and p['method']=='uniform')
            rr.append([r['count'],shapes[r['shape']],f"{u['control_model_com_error_m']*1000:.2f}",'规划候选拒绝' if u['motion_start_sec'] is None else '姿态及加速度超限','通过'])
        table('表 3 均匀外形与辨识模型配对结果',['机数','外形','均匀模型质心误差\n(mm)','均匀模型结果','辨识模型结果'],rr,[1.0,2.,4.3,5.,4.3])
    elif key=='ablation':
        rr=[];allrows=json.loads((ROOT/'supplementary/review_20261027/TRIAL_METRICS.json').read_text(encoding='utf-8'))
        for n in (3,4):
            for shape in ('rectangular','irregular'):
                cells=[]
                for method in ('mass_matched','mass_only','mass_com','estimated','oracle'):
                    r=next(r for r in allrows if r['count']==n and r['shape']==shape and r['method']==method and r['seed']==2501 and not r['retry'])
                    cells.append((f"{r['peak_attitude_deg']:.2f}"+('*' if not r['passed'] else '')) if r['completed'] else ('A' if r['reason']=='image_tracking_stale' else 'R'))
                rr.append([n,shapes[shape]]+cells)
        table('表 4 五组公平消融。数值为完成运输的姿态峰值（°）；R 为候选拒绝，A 为起运前图像年龄中止，* 为执行超限。每格为首次尝试。',['机数','外形','真质量\n均匀模型','仅质量','质量\n与质心','完整\n辨识','真实\n惯性'],rr,[.9,1.7,3.1,2.2,3.1,2.8,2.8])
    elif key=='course':
        rr=[]
        for i,c in enumerate(E['courses']):
            r=c['summary'];g=c['geometry'];rr.append(['中等转姿' if i==0 else '大幅转姿',f"{r['planned_motion_duration_sec']:.0f}",f"{r['max_position_error_m']*100:.2f}",f"{r['max_attitude_error_deg']:.2f}",f"{r['min_tension_n']:.2f}–{r['max_tension_n']:.2f}",f"{g['minimum_aircraft_separation_m']:.3f}",f"{r['initial_planning_wall_sec']:.2f}"])
        table('表 5 叉形载荷课程结果。两课程首门宽度及参考转姿幅度不同。',['课程','运输\n(s)','位置峰值\n(cm)','姿态峰值\n(°)','张力\n(N)','最小机距\n(m)','规划\n(s)'],rr,[2.2,1.5,2.5,2.5,3.2,2.6,2.1])

def references():
    text=(M/'references.tex').read_text(encoding='utf-8')
    parts=re.split(r'\\bibitem\{[^}]+\}',text)[1:]
    for i,item in enumerate(parts,1):
        item=item.split(r'\end{thebibliography}')[0].strip()
        item=re.sub(r'\\doi\{([^}]+)\}',r'https://doi.org/\1',item)
        item=item.replace('--','–').replace('\\&','&')
        p=paragraph(f'[{i}] {item}');p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3)
        for r in p.runs:r.font.size=Pt(8.2)

source=(M/'article_zh.md').read_text(encoding='utf-8').replace('꜀','c')
first_figure=re.search(r'!\[图 1 .*?\]\(fig01_overview.png\)',source,re.S)[0]
source=source.replace(first_figure+'\n\n','',1).replace('![图 2 ',first_figure+'\n\n![图 2 ',1)
for block in re.split(r'\n\s*\n',source):
    block=block.strip()
    if not block:continue
    if block.startswith('# '):
        p=paragraph(block[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        author_block()
    elif block.startswith('### '):doc.add_heading(block[4:],2)
    elif block.startswith('## '):doc.add_heading(block[3:],1)
    elif block.startswith('EQ|'):equation(block[3:])
    elif block.startswith('TABLE|'):add_table(block[6:])
    elif block=='REFERENCES':references()
    elif block.startswith('!['):
        match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',block,re.S)
        p=paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        width=15.8 if 'fig01_' not in match[2] else 14.8
        p.add_run().add_picture(str(M/'figures'/match[2]),width=Cm(width))
        p=paragraph(match[1],'Caption');p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p=paragraph(block)
        if block=='基于视觉与张力感知':p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif not block.startswith('关键词'):p.paragraph_format.first_line_indent=Cm(.74)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='面向障碍穿越的多无人机协同吊运载荷变姿与编队重构'
doc.core_properties.author='Fangwei Peng; Juntong Qi; Yan Peng; Mingming Wang; Chong Wu'
doc.core_properties.subject='Vision and tension informed cooperative aerial transport'
for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
output=ROOT/'ForestSling_AST_Chinese_Manuscript.docx';doc.save(output)
print(output)
